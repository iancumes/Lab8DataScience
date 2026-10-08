#!/usr/bin/env python3
"""Genera el informe Word a partir de los resultados guardados del laboratorio.

Dependencia opcional: python-docx==1.2.0 (requirements-report.txt).
"""
import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'output/docx/Laboratorio_8_DuckDB.docx'
STUDENTS = ['Javier Valladares, 23045', 'Ian Cumes, 23236', 'Javier Cifuentes, 23079']


def rows(name, stage='final'):
    with (ROOT / 'docs/results' / stage / f'{name}.csv').open() as stream:
        return list(csv.DictReader(stream))


def number(value, digits=0):
    return f'{float(value):,.{digits}f}'


doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.8)
section.footer_distance = Inches(0.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.08
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Title'].font.size = Pt(24)
doc.styles['Title'].font.bold = True
doc.styles['Title'].font.underline = False
doc.styles['Subtitle'].font.italic = False
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
for name, size in [('Heading 1', 15), ('Heading 2', 11)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.space_before = Pt(10)
    doc.styles[name].paragraph_format.keep_with_next = True
doc.styles['Caption'].font.size = Pt(9)
doc.styles['Caption'].font.italic = False
doc.core_properties.title = 'Laboratorio 8 DuckDB'
doc.core_properties.author = '; '.join(STUDENTS)
doc.core_properties.subject = 'Análisis de viajes de taxis de Nueva York'
doc.core_properties.comments = ''
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)


def p(text, style=None):
    return doc.add_paragraph(text, style)


next_page = False


def h(text, level=1):
    global next_page
    paragraph = doc.add_heading(text, level)
    if next_page:
        paragraph.paragraph_format.page_break_before = True
        next_page = False
    return paragraph


def page():
    global next_page
    next_page = True


def table(headers, data, widths, size=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement(f'w:{edge}')
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            el.set(qn(f'w:{key}'), value)
        borders.append(el)
    pr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge in ['top', 'left', 'bottom', 'right']:
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:w'), '75')
        el.set(qn('w:type'), 'dxa')
        margins.append(el)
    pr.append(margins)
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for record in [headers] + list(data):
        row = t.rows[0] if record is headers else t.add_row()
        for i, (cell, value, width) in enumerate(zip(row.cells, record, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(value)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                run.font.size = Pt(size)
                run.font.bold = record is headers
            if record is headers:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'EDEDED')
                cell._tc.get_or_add_tcPr().append(shading)
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
    p('').paragraph_format.space_after = Pt(0)
    return t


def code(text):
    paragraph = p(text)
    paragraph.paragraph_format.line_spacing = 1.0
    for run in paragraph.runs:
        run.font.name = 'Courier New'
        run.font.size = Pt(8.5)


def picture(name, width=6.9):
    paragraph = p('')
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run().add_picture(str(ROOT / 'docs/figures' / name), width=Inches(width))


verification = json.loads((ROOT / 'docs/verification.json').read_text())
quality = rows('03_quality')
monthly = rows('04_monthly')
common = rows('08_common_months')
initial_quality = rows('03_quality', 'initial_2026')
initial_monthly = rows('04_monthly', 'initial_2026')
initial_rows = sum(int(r['rows']) for r in initial_quality)
initial_excluded = sum(int(r['excluded']) for r in initial_quality)
total_rows = sum(int(r['rows']) for r in quality)
excluded = sum(int(r['excluded']) for r in quality)
taxi_names = {'yellow': 'Amarillo', 'green': 'Verde'}

p('Laboratorio 8 DuckDB', 'Title')
p('Análisis de viajes de taxis de Nueva York', 'Subtitle')
p('CC3084 Ciencia de Datos\nUniversidad del Valle de Guatemala\nCiclo 2 de 2026')
for student in STUDENTS:
    p(student)
p('8 de octubre de 2026')
p('Repositorio: https://github.com/iancumes/Lab8DataScience')
p('Analizamos archivos Parquet de taxis amarillos y verdes de 2024, 2025 y 2026. '
  'El conjunto contiene 64 archivos y 121,184,384 registros. El año 2026 llega hasta agosto, '
  'por lo que comparamos los años usando los mismos ocho meses. Los taxis amarillos tienen '
  'mucho más volumen. El benchmark también muestra que materializar los datos no acelera todas las consultas.')
h('Ejercicio 1 Ambiente de trabajo')
p('Usamos Docker Compose con dos servicios. JupyterLab permite ejecutar el notebook y Metabase '
  'muestra el tablero. DuckDB consulta los Parquet y genera tablas pequeñas de resultados. '
  'La lectura de datos, las pruebas de descarga y la verificación final también se comprobaron dentro de Docker.')
table(['Componente', 'Versión o acceso'], [
    ['Python del contenedor', '3.11.14'], ['DuckDB', '1.5.5'],
    ['Pandas y PyArrow', '3.0.6 y 25.0.1'], ['JupyterLab', '4.6.4, puerto 8888'],
    ['Metabase y driver DuckDB', 'v0.63.19 y 1.5.5.0, puerto 3000']], [2.8, 4.1])
p('Los datos quedan en data/raw y las bases generadas en data/processed. Ambas carpetas se '
  'excluyen de Git. El código, las consultas SQL y los resultados resumidos sí se versionan. '
  'Docker y las versiones fijadas ayudan a reproducir el ambiente. Los manifiestos identifican los archivos usados.')

page()
h('Ejercicio 2 Descarga de datos')
p('Partimos del repositorio de menene/duckdb y ampliamos el descargador para recibir años y meses como '
  'parámetros. Se consultan los enlaces del catálogo oficial de NYC TLC. Cada archivo se descarga '
  'primero como temporal, se revisa su tamaño y se comprueba que PyArrow pueda leer sus metadatos. '
  'Luego se guarda con su nombre final. También registramos filas, bytes y hash SHA-256.')
p('Si un archivo ya existe, se valida y se conserva. Un archivo corrupto se reporta como error y '
  'no se sobrescribe. Así evitamos volver a descargar datos que ya estaban completos.')
table(['Etapa', 'Archivos', 'Filas', 'Nuevos', 'Existentes'], [
    [label, number(v['files']), number(v['rows']), v['downloaded'], v['existing']]
    for key, label in [('initial_2026', '2026'), ('expanded_2024_2026', '2024 y 2026'),
                       ('final', 'Tres años'), ('rerun', 'Segunda ejecución')]
    for v in [verification['stages'][key]]], [1.65, 0.8, 2.05, 1.1, 1.3])
p('La segunda ejecución encontró los 64 archivos existentes y no descargó ninguno. En esta entrega, '
  '2024 y 2025 tienen doce meses por tipo de taxi y 2026 tiene ocho. El catálogo puede publicar meses adicionales después del corte.')
h('Ejercicio 3 Exploración directa de Parquet')
p('Consultamos los archivos con read_parquet, sin cargarlos completos en Pandas. Revisamos filas, '
  'columnas, tipos y diez registros de muestra por tipo de taxi. La muestra sirve para revisar la estructura, '
  'pero no se usa para estimar los indicadores del conjunto completo.')
table(['Taxi', 'Año', 'Archivos', 'Filas originales'], [
    [taxi_names[r['taxi']], r['source_year'], r['files'], number(r['rows'])]
    for r in rows('01_inventory')], [1.7, 1.0, 1.2, 3.0])
p('Normalizamos las fechas tpep de yellow y lpep de green como pickup y dropoff. '
  'La vista común tiene 19 columnas. union_by_name permite unir archivos cuyos esquemas cambian. '
  'Los esquemas originales se conservan en los CSV de resultados.')
code("SELECT COUNT(*) AS filas\nFROM read_parquet(\n  'data/raw/yellow/2026/*.parquet',\n  union_by_name = true\n);")

page()
h('Ejercicio 4 Hallazgos iniciales de 2026')
p('En la primera etapa se analizaron 30,040,469 registros de enero a agosto de 2026. '
  'Después de aplicar las reglas de calidad quedaron 28,219,350 viajes amarillos y 317,814 verdes. '
  'La diferencia de volumen es de aproximadamente 89 viajes amarillos por cada viaje verde.')
table(['Indicador de 2026', 'Amarillo', 'Verde'], [
    ['Viajes válidos', '28,219,350', '317,814'],
    ['Importe medio por viaje en USD', '30.24', '25.35'],
    ['Hora con más viajes', '18:00', '17:00']], [3.5, 1.7, 1.7])
p(f'Las reglas excluyeron {number(initial_excluded)} filas, equivalentes al '
  f'{100 * initial_excluded / initial_rows:.2f}% de la etapa inicial. Encontramos 244 fechas '
  'fuera del mes y año del archivo o sin fecha, además de 7,765,463 registros sin passenger_count. '
  'Las fechas antiguas, como 2001 y 2008, muestran por qué no conviene agrupar solamente por la fecha registrada sin revisar su coherencia.')
h('Reglas usadas en los indicadores', 2)
p('Un viaje entra en la población analítica si su fecha de inicio coincide con el mes y año del archivo, '
  'su distancia es mayor que cero y no supera 100 millas, su duración es mayor que cero y no supera '
  '180 minutos, y su tarifa e importe total son positivos. Son criterios de este análisis. Algunos '
  'viajes excluidos pueden corresponder a ajustes, reembolsos o recorridos largos reales.')
p('Los archivos originales y las filas marcadas se conservan. Los motivos de exclusión pueden '
  'solaparse, por lo que sus conteos no se suman directamente. No imputamos pasajeros ni eliminamos '
  'supuestos duplicados, porque no hay un identificador universal que permita confirmarlos.')
h('Ejercicio 5 Incorporación de 2024')
p('Agregamos los 24 archivos de 2024 y conservamos los 16 de 2026. El conjunto pasó a 40 archivos '
  'y 71,870,407 registros. Las consultas se volvieron a ejecutar con la misma lógica. Solo cambió '
  'el parámetro de años, de 2026 a 2024 y 2026.')
p('El manifiesto expanded_2024_2026 registra los archivos anteriores como existing. La verificación '
  'compara sus hashes, tamaños y filas para comprobar que se conservaron. Esta etapa permite revisar '
  'el comportamiento del análisis cuando crece el volumen de datos.')
p('El campo cbd_congestion_fee no existe en los archivos de 2024. Se representa como NULL para '
  'distinguir la ausencia de la columna de un recargo igual a cero. El mismo criterio se mantiene '
  'al incorporar 2025. Los resultados y las consultas de cada etapa quedan en docs/results.')

page()
h('Ejercicio 6 Parquet y tablas DuckDB')
p('Comparamos cuatro consultas en cuatro escalas. Ambas rutas usan las mismas filas y las mismas '
  '19 columnas normalizadas. Una lee Parquet y la otra una tabla DuckDB materializada. '
  'Calentamos ambas rutas y realizamos tres repeticiones alternando el orden. El tiempo incluye '
  'ejecutar y recuperar el resultado. Las 96 mediciones verifican que ambas rutas devuelven resultados equivalentes.')
p('El benchmark se ejecutó en macOS con Python 3.13.0, DuckDB 1.5.5, cuatro hilos y un límite de '
  'memoria de 3 GB. Se usó caché caliente y no se vació la caché del sistema operativo. '
  'La descarga queda fuera de la medición. La creación de la tabla y CHECKPOINT se midieron aparte.')
with (ROOT / 'docs/benchmarks/materialization.csv').open() as stream:
    mats = list(csv.DictReader(stream))
scale_names = {'2026_one_month': 'Un mes de 2026', '2026': '2026',
               '2024_2026': '2024 y 2026', '2024_2025_2026': 'Tres años'}
table(['Escala', 'Filas', 'Crear tabla en s', 'Base en GB'], [
    [scale_names[r['scale']], number(r['rows']), number(r['materialization_seconds'], 3),
     number(float(r['database_bytes']) / 1e9, 3)] for r in mats], [1.7, 2.1, 1.65, 1.45])
p('La siguiente tabla muestra las medianas en segundos. Distribución calcula cuantiles de importes, '
  'horario agrupa por hora, mensual agrupa por mes y zonas obtiene zonas principales con un filtro.')
with (ROOT / 'docs/benchmarks/summary.csv').open() as stream:
    summary = list(csv.DictReader(stream))
bench = {(r['scale'], r['query'], r['mode']): float(r['median']) for r in summary}
query_names = {'distribution': 'Distribución', 'hourly': 'Horario', 'monthly': 'Mensual', 'top_zones': 'Zonas'}
table(['Escala', 'Consulta', 'Parquet en s', 'Tabla en s'], [
    [scale_names[scale], query_names[query], f'{bench[scale, query, "parquet"]:.4f}',
     f'{bench[scale, query, "duckdb"]:.4f}']
    for scale in scale_names for query in query_names], [1.9, 1.8, 1.6, 1.6], size=9)
p('En los tres años, la tabla mejora la consulta por hora de 1.1341 a 0.6592 segundos. '
  'La consulta mensual favorece Parquet, con 0.5018 frente a 0.6331 segundos. La suma de las '
  'cuatro consultas no ahorra tiempo al materializar en esta medición, y crear la tabla cuesta '
  '30.891 segundos adicionales. Para esta carga completa no observamos un punto de amortización positivo.')

page()
h('Ejercicio 7 Preguntas e indicadores')
p('Definimos trece preguntas para orientar el análisis. Las primeras diez se presentan en el tablero '
  'Metabase mediante diez tarjetas SQL. Las restantes complementan la revisión de patrones y valores extremos.')
questions = [
    ['1 Demanda mensual', 'Viajes válidos por mes', 'Comparar volumen y estacionalidad.'],
    ['2 Importe mensual', 'Suma de total_amount', 'Medir actividad registrada en USD.'],
    ['3 Importe por viaje', 'Media de total_amount', 'Comparar el pago medio registrado.'],
    ['4 Distancia típica', 'Mediana en millas', 'Resumir distancias con menor efecto de extremos.'],
    ['5 Duración típica', 'Mediana en minutos', 'Comparar el tiempo de los recorridos.'],
    ['6 Horario de actividad', 'Viajes por hora', 'Identificar horas con más demanda.'],
    ['7 Mezcla de pagos', 'Porcentaje por código', 'Revisar cambios de registro y medio de pago.'],
    ['8 Propinas en tarjeta', 'Propinas sobre tarifa', 'Medir propinas que sí están registradas.'],
    ['9 Calidad de datos', 'Porcentaje excluido', 'Mostrar cuánto cambia la población analizada.'],
    ['10 Evolución anual', 'Viajes en meses comunes', 'Comparar años con la misma cobertura.'],
    ['11 Día de semana', 'Viajes por día observado', 'Ajustar por la cantidad de cada día de semana.'],
    ['12 Distribuciones', 'Bandas y cuantiles', 'Revisar colas y cambios de composición.'],
    ['13 Valores extremos', 'Veinte importes mayores', 'Revisar casos sin corregirlos automáticamente.'],
]
table(['Pregunta', 'Indicador', 'Motivo'], questions, [2.05, 2.05, 2.8], size=10)
p('El tablero usa un archivo DuckDB con agregados, separado de la base del benchmark y abierto '
  'en modo de solo lectura. Así las tarjetas consultan resultados pequeños. El script '
  'setup_metabase.py crea o actualiza las diez visualizaciones y comprueba sus consultas. '
  'El registro de tarjetas y SQL queda en docs/dashboard.json.')
p('El tablero verificado está disponible localmente en http://localhost:3000/dashboard/2. '
  'Los identificadores pueden cambiar en otra instalación. La captura del tablero real se encuentra '
  'en docs/tablero.md y las consultas están en sql/dashboard. El color dorado representa yellow y '
  'el verde representa green. Se distinguen sus escalas para poder observar ambos servicios.')
p('El importe total no equivale a beneficio del conductor. Las propinas en efectivo no están '
  'registradas. El código de pago 0 corresponde a viajes Flex Fare y no identifica un instrumento '
  'de pago, por lo que una caída del código 1 no demuestra un aumento del efectivo.')

page()
h('Resultados de los indicadores')
p('Los gráficos resumen los tres años. Los indicadores mensuales usan los viajes que cumplen las '
  'reglas de calidad. El gráfico de calidad usa todas las filas originales. El año 2026 es parcial.')
picture('indicadores.png')
p('Figura 1 Indicadores calculados con DuckDB y graficados con Matplotlib', 'Caption')
p('Para dar una referencia concreta, estos son los resultados de agosto de 2026, el último mes de la descarga.')
august = {r['taxi']: r for r in monthly if r['year'] == '2026' and r['month'] == '8'}
table(['Indicador de agosto de 2026', 'Amarillo', 'Verde'], [
    [label, number(august['yellow'][key], digits), number(august['green'][key], digits)]
    for key, label, digits in [
        ('trips', 'Viajes válidos', 0), ('total_usd', 'Importe registrado en USD', 2),
        ('avg_total_usd', 'Importe medio en USD', 2), ('median_miles', 'Distancia mediana en millas', 2),
        ('median_minutes', 'Duración mediana en minutos', 2), ('credit_pct', 'Código tarjeta en porcentaje', 2),
        ('card_tip_pct', 'Propina sobre tarifa en tarjeta en porcentaje', 2)]], [3.5, 1.7, 1.7], size=9)
p('En ese mes, los taxis amarillos registraron mucho más volumen e importe total. Los verdes '
  'tuvieron una distancia mediana algo mayor y una duración mediana menor. Estas diferencias '
  'describen los viajes observados y pueden reflejar recorridos y zonas distintos.')

page()
h('Ejercicio 8 Incorporación de 2025 y evolución')
p('Añadimos 24 archivos de 2025 sin descargar nuevamente 2024 ni 2026. El conjunto final tiene '
  '64 archivos y 121,184,384 registros. Se ejecutaron las mismas consultas y se actualizó el tablero. '
  'Para comparar años usamos enero a agosto, porque son los meses presentes en los tres años.')
table(['Taxi', 'Año', 'Viajes válidos', 'Media USD', 'Minutos medios', 'Tarjeta %'], [
    [taxi_names[r['taxi']], r['year'], number(r['trips']), number(r['avg_total_usd'], 2),
     number(r['avg_minutes'], 2), number(r['credit_pct'], 2)]
    for taxi in ['yellow', 'green'] for r in common if r['taxi'] == taxi],
    [1.0, 0.65, 1.65, 1.15, 1.25, 1.2], size=9)
p('La demanda amarilla aumentó 12.52% entre 2024 y 2025 y disminuyó 1.62% entre 2025 y 2026. '
  'La demanda verde cayó 10.23% y luego 15.05%. Estas variaciones corresponden a los mismos '
  'ocho meses, sin extrapolar al resto de 2026.')
p('El importe medio amarillo pasó de USD 28.32 en 2024 a USD 30.24 en 2026, un aumento de 6.77%. '
  'En green pasó de USD 23.74 a USD 25.35, un aumento de 6.78%. Los importes son nominales y '
  'no se ajustaron por inflación. Los cambios de recorrido, cargos y composición también pueden influir.')
picture('perfil_horario.png')
p('Figura 2 Distribución de viajes por hora según tipo de taxi y año', 'Caption')
p('El perfil horario se expresa como porcentaje para comparar años con distinta cobertura. '
  'Al agregar todo el período, la hora con más viajes es 18:00 para amarillo y 17:00 para verde. '
  'Los patrones por hora son descriptivos y no prueban una causa del cambio de demanda.')
p(f'En el conjunto final se excluyen {number(excluded)} filas, el '
  f'{100 * excluded / total_rows:.2f}% del total. Hay 1,290 fechas incoherentes o ausentes y '
  '23,542,797 filas sin número de pasajeros. Los resultados principales corresponden a la '
  'población filtrada. cbd_congestion_fee aparece desde 2025 y se conserva como NULL en 2024. '
  'El cambio del ticket no permite atribuir un efecto causal a ese recargo.')

page()
h('Ejercicio 9 Discusión')
h('1 Características útiles de DuckDB', 2)
p('La lectura directa de Parquet permitió empezar el análisis sin una importación previa. '
  'Usamos agrupaciones, ventanas y cuantiles con SQL. union_by_name ayudó a combinar esquemas '
  'distintos. DuckDB se ejecuta dentro del proceso y evita administrar un servidor de base de datos para estas consultas.')
h('2 Ventajas y límites de Parquet', 2)
p('Parquet conserva los datos originales y evita mantener otra copia completa. DuckDB puede leer '
  'solo las columnas necesarias y aprovechar filtros. Hay que administrar rutas, particiones y '
  'cambios de esquema. Las consultas repetidas pueden volver a decodificar columnas. COUNT puede '
  'usar metadatos, por lo que el benchmark incluye operaciones más exigentes.')
h('3 Ventajas y límites de materializar', 2)
p('Una tabla ofrece una copia normalizada estable y puede acelerar consultas repetidas. En este '
  'equipo mejoró el perfil horario, pero no la consulta mensual. La creación cuesta tiempo y '
  'espacio, y la tabla necesita actualizarse cuando llegan archivos nuevos. Para Metabase '
  'usamos otro archivo con agregados y conexión de solo lectura para evitar conflictos con el análisis.')
h('4 Comparación con Pandas', 2)
p('DuckDB filtra y agrega antes de enviar resultados a un DataFrame. Esto evita cargar los '
  '121 millones de registros completos en Pandas. También puede usar disco temporal cuando '
  'la memoria no alcanza. En este laboratorio Pandas recibió tablas pequeñas. No medimos '
  'una comparación de velocidad contra Pandas, así que no podemos afirmar que siempre sea más rápido.')
h('5 Incorporación de archivos nuevos', 2)
p('Los años y meses son parámetros. El descargador descubre enlaces oficiales, valida los '
  'archivos existentes y descarga los que faltan. Las consultas no contienen una lista fija '
  'de meses. Al ejecutar el análisis se reconstruyen las vistas. Las tablas materializadas '
  'y el archivo del tablero se regeneran de forma explícita.')
h('6 Uso en producción', 2)
p('Automatizaríamos la consulta del catálogo, las descargas, la validación y la ejecución del '
  'análisis. Guardaríamos manifiestos y versiones de entradas, con alertas ante cambios de '
  'esquema o conteos inusuales. El refresco del tablero debe coordinarse con sus lectores. '
  'También se necesitan reintentos, registro de errores y controles de resultados.')
h('7 Reproducibilidad', 2)
p('La entrega incluye Docker, versiones directas fijadas, SQL, reglas de calidad, hashes, '
  'manifiestos por etapa y el entorno del benchmark. Las rutas dependen de la carpeta del '
  'proyecto. Los datos completos y las credenciales quedan fuera de Git. Las imágenes base, '
  'dependencias transitivas y archivos remotos pueden cambiar, por lo que no prometemos una reproducción idéntica byte por byte.')
h('8 Aprendizaje al aumentar la escala', 2)
p('Al trabajar con millones de filas fue necesario separar datos y código, controlar el espacio '
  'en disco y reducir las columnas antes de enviar resultados a Python. También se hicieron '
  'visibles fechas anómalas y cambios de esquema. El resultado de una muestra pequeña o de '
  'una sola consulta no basta para decidir si conviene materializar todo el conjunto.')

page()
h('Verificación y reproducción')
p('La verificación registrada terminó correctamente. Comprobó los 64 hashes, la conservación '
  'de archivos entre etapas, las 96 mediciones equivalentes, las diez tarjetas con datos y el '
  'notebook ejecutado. Las cuatro pruebas del descargador cubren archivos existentes, '
  'archivos corruptos, Parquet vacíos y descargas interrumpidas.')
p('El README contiene todos los comandos. El flujo principal desde una copia del repositorio es el siguiente. '
  'Las descargas completas requieren varios GB libres. Antes de configurar Metabase hay que esperar '
  'a que su comprobación de salud responda correctamente.')
code('git clone https://github.com/iancumes/Lab8DataScience.git\n'
     'cd Lab8DataScience\ntouch .env\ndocker compose up -d --build\n'
     'docker compose exec lab python scripts/run_pipeline.py\n'
     'docker compose exec lab python scripts/create_notebook.py\n'
     'curl -f http://localhost:3000/api/health\n'
     'docker compose exec lab python scripts/setup_metabase.py \\\n'
     '  --url http://metabase:3000\n'
     'docker compose exec lab python scripts/download_data.py \\\n'
     '  --years 2024 2025 2026 --manifest docs/manifests/rerun.json\n'
     'docker compose exec lab python scripts/verify.py')
p('La configuración inicial de Metabase guarda el acceso local en .env. Ese archivo está '
  'excluido de Git. El tablero se consulta en el puerto 3000 y JupyterLab en el puerto 8888. '
  'Los CSV y el notebook guardados permiten revisar esta ejecución sin descargar de inmediato todos los datos.')
h('Referencias')
for text in [
    'Repositorio base: https://github.com/menene/duckdb',
    'NYC TLC, catálogo y notas de calidad: https://www.nyc.gov/site/tlc/about/tlc-trip-record-data.page',
    'Diccionario yellow: https://www.nyc.gov/assets/tlc/downloads/pdf/data_dictionary_trip_records_yellow.pdf',
    'Diccionario green: https://www.nyc.gov/assets/tlc/downloads/pdf/data_dictionary_trip_records_green.pdf',
    'DuckDB, Parquet: https://duckdb.org/docs/stable/data/parquet/overview',
    'DuckDB, concurrencia: https://duckdb.org/docs/stable/connect/concurrency',
    'Driver para Metabase: https://github.com/motherduckdb/metabase_duckdb_driver',
    'Enunciado del curso: Lab_8___DuckDB.pdf. Resultados de la ejecución del 8 de octubre de 2026 en docs/results, docs/manifests y docs/benchmarks.',
]:
    p(text)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
